#!/usr/bin/env python3
"""Export the SQLite snapshot to reviewable XLSX, JSONL and a compact receipt."""
import argparse
from datetime import datetime
import json
from pathlib import Path
import sqlite3
from zoneinfo import ZoneInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parents[1]
HEADERS_KO = {
 'equipment_id':'장비ID', 'register_no':'원본 No.', 'name':'장비명',
 'manufacturer':'제조사 표준표기', 'model_name':'모델명', 'serial_number':'일련번호',
 'domain':'대분류(추정)', 'category':'장비종류(추정)', 'specification':'장비사양',
 'calibration_specification':'교정사양', 'note':'비고', 'firmware':'펌웨어·SW(정리값)',
 'classification_status':'분류 검토상태', 'classification_rule':'적용 분류규칙',
 'automation_research_hint':'자동화 조사방향(미검증)', 'automation_status':'자동화 상태',
 'lifecycle_status':'현재 운용상태', 'source_status_hint':'비고 상태단서',
 'location':'설치위치(미확인)', 'owner':'담당자(미확인)', 'calibration_due_date':'교정 만료일(미확인)',
 'manufacturer_id':'제조사ID', 'model_id':'모델ID', 'manufacturer_raw':'제조사 원문',
 'name_raw':'장비명 원문','model_raw':'모델 원문','serial_raw':'일련번호 원문',
 'specification_raw':'장비사양 원문','calibration_specification_raw':'교정사양 원문',
 'note_raw':'비고 원문','firmware_raw':'펌웨어 원문',
 'sheet_name':'원본시트', 'source_row':'원본행', 'original_filename':'원본파일', 'source_sha256':'원본 SHA-256',
 'canonical_name':'제조사 표준표기','equipment_records':'등록건수','model_count':'모델수',
 'source_label_count':'원문 표기수','verification_status':'확인상태',
 'raw_label':'제조사 원문','normalization_method':'표기 정리방법','review_required':'검토필요',
 'suggested_canonical_name':'제조사 정정후보(미적용)','rationale':'정리 근거',
 'suggested_categories':'장비종류 후보','research_hints':'자동화 조사방향(미검증)',
 'research_status':'조사상태','remote_control_supported':'원격제어 지원(미확인)',
 'manual_url':'사용설명서 URL','programming_manual_url':'프로그래밍 설명서 URL',
 'sdk_url':'SDK URL','driver_key':'드라이버 식별자','notes':'조사메모','verified_at':'확인일시',
 'category_id':'분류ID','domain_ko':'대분류','category_ko':'장비종류',
 'rule_pattern':'이름 매칭규칙','issue_id':'확인항목ID','field_name':'대상필드',
 'issue_type':'확인유형','severity':'중요도','detail':'확인내용','review_status':'처리상태',
 'duplicate_group_id':'중복후보 그룹ID','match_type':'일치기준','member_count':'그룹 등록건수',
 'resolution_status':'검토상태','identifier_type':'식별번호 유형','identifier_value':'식별번호',
 'source_field':'추출필드','source_excerpt':'추출 원문','identifier_id':'식별번호ID',
}
LABELS_KO = {
 'suggested':'이름 기준 1차 분류', 'unclassified':'미분류·확인 필요', 'reviewed':'검토 완료',
 'not_assessed':'미확인', 'source_only':'원본 기반·외부 미확인', 'unknown':'미확인',
 'no_status_evidence':'원본에 상태 단서 없음', 'status_history_review':'상태 이력 확인 필요',
 'retired_reported':'비고에 폐기 명시', 'out_of_service_reported':'비고에 사용중지 명시',
 'source_extracted':'원문에서 추출', 'legacy_registration':'과거 등록번호',
 'format_only':'대소문자·공백 등 정리', 'explicit_alias':'명시적 별칭 매핑',
 'suspected_typo_unmerged':'오탈자 의심·병합 안 함', 'missing':'결측',
 'non_manufacturer':'제조사 아닌 표기', 'multiple_manufacturers':'복수 제조사 기재',
 'open':'미검토', 'unreviewed':'미검토', 'resolved':'확인 완료',
 'high':'높음', 'medium':'보통', 'low':'낮음',
 'exact_source_fields':'No. 제외 원문 8개 항목 일치',
 'manufacturer_model_serial':'제조사·모델·일련번호 일치',
 'serial_only':'일련번호만 일치(약한 근거)',
 'missing_identifier':'식별정보 누락·미지정',
 'manufacturer_alias_review':'제조사 오탈자 후보 확인',
 'model_catalog_unresolved':'제조사·모델 조합 미확정',
 'lifecycle_note_review':'현재 운용상태 확인 필요',
}
ENUM_COLUMNS = {'classification_status','automation_status','lifecycle_status','source_status_hint',
                'verification_status','normalization_method','research_status','review_status',
                'severity','match_type','resolution_status','identifier_type','issue_type'}


def fetch(con, sql):
    return [dict(r) for r in con.execute(sql)]


def summary(con):
    scalar = lambda q: con.execute(q).fetchone()[0]
    return {
      'generated_at_kst':datetime.now(ZoneInfo('Asia/Seoul')).isoformat(),
      'grain':'원본 관리대장의 등록 1행; 실물 장비 고유 대수와 다를 수 있음',
      'source':fetch(con,'SELECT original_filename,sha256,record_count,sheet_count FROM source_documents')[0],
      'equipment_records':scalar('SELECT count(*) FROM equipment'),
      'raw_manufacturer_labels':scalar('SELECT count(*) FROM manufacturer_aliases'),
      'manufacturer_groups':scalar('SELECT count(*) FROM manufacturers'),
      'manufacturer_groups_are_verified_companies':False,
      'models':scalar('SELECT count(*) FROM models'),
      'equipment_without_model_master':scalar('SELECT count(*) FROM equipment WHERE model_id IS NULL'),
      'equipment_without_manufacturer':scalar('SELECT count(*) FROM equipment WHERE manufacturer_id IS NULL'),
      'equipment_without_model':scalar('SELECT count(*) FROM equipment WHERE model_name IS NULL'),
      'equipment_without_serial':scalar('SELECT count(*) FROM equipment WHERE serial_number IS NULL'),
      'equipment_without_name':scalar('SELECT count(*) FROM equipment WHERE name IS NULL'),
      'unclassified_records':scalar("SELECT count(*) FROM equipment WHERE classification_status='unclassified'"),
      'legacy_identifiers':scalar('SELECT count(*) FROM equipment_identifiers'),
      'category_count':scalar('SELECT count(*) FROM categories'),
      'category_counts':fetch(con,'SELECT c.category_id,c.domain_ko,c.category_ko,count(e.equipment_id) equipment_records FROM categories c LEFT JOIN equipment e USING(category_id) GROUP BY c.category_id ORDER BY equipment_records DESC'),
      'quality_counts':fetch(con,'SELECT issue_type,field_name,count(*) affected_records FROM data_quality_issues GROUP BY 1,2 ORDER BY 3 DESC'),
      'status_hints':fetch(con,'SELECT source_status_hint,count(*) equipment_records FROM equipment GROUP BY 1'),
      'duplicates':fetch(con,'SELECT match_type,count(*) candidate_groups,sum(member_count) affected_records FROM duplicate_groups GROUP BY 1'),
      'verified_connections':scalar("SELECT count(*) FROM equipment_connections WHERE verification_status='bench_verified'"),
      'supported_model_interfaces':scalar('SELECT count(*) FROM model_interfaces'),
      'caveats':[
        f"원본 {scalar('SELECT count(*) FROM equipment'):,}행을 삭제·병합하지 않음. 등록건수는 운용 가능 실물 장비 대수가 아님.",
        '제조사 표준표기는 표기 그룹. 법인 확인·인수합병 통합·공식 매뉴얼 검증을 수행하지 않음.',
        '장비종류는 이름 기반 추정. 미분류와 분류 규칙을 함께 제공.',
        'SCPI/LAN/USB/GPIB/SDK 지원, IP·주소, 현행 교정 유효기간은 미확인.',
        '원본 No.는 관리대장 순번. 장비ID는 최초 스냅샷 식별자이며 실물 자산번호로 단정하지 않음.',
        '펌웨어 원문 항목에는 소프트웨어명도 포함될 수 있음.',
        '원본 파일의 재고 기준일은 확인되지 않음. 생성일시는 DB·내보내기 작업 일시.',
      ],
    }


def write_sheet(wb, title, rows, keys=None):
    ws=wb.create_sheet(title)
    if keys is None:
        keys=list(rows[0]) if rows else []
    headers=[HEADERS_KO.get(k,k) for k in keys]
    ws.append(headers)
    for row_number,row in enumerate(rows,2):
        values=[]
        for key in keys:
            value=row.get(key)
            if key in ENUM_COLUMNS:
                value=LABELS_KO.get(value,value)
            values.append(value)
        ws.append(values)
        # Treat every source string as data, including strings beginning with '='.
        for column in range(1,len(keys)+1):
            cell=ws.cell(row_number,column)
            if isinstance(cell.value,str):
                cell.data_type='s'
    ws.freeze_panes='D2' if len(keys)>5 else 'A2'
    ws.sheet_view.showGridLines=False
    for cell in ws[1]:
        cell.fill=PatternFill('solid',fgColor='16324F')
        cell.font=Font(color='FFFFFF',bold=True,size=11)
        cell.alignment=Alignment(vertical='center',wrap_text=True)
    ws.row_dimensions[1].height=32
    for column,key in enumerate(keys,1):
        width=22
        if key in ('register_no','source_row','equipment_records','model_count','source_label_count'):
            width=13
        elif key in ('name','name_raw','specification','calibration_specification','note','note_raw','detail','rationale','source_excerpt'):
            width=48
        elif key.endswith('_id'):
            width=26
        elif 'url' in key or key in ('source_sha256','classification_rule','rule_pattern'):
            width=38
        ws.column_dimensions[get_column_letter(column)].width=width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment=Alignment(vertical='top',wrap_text=True)
            cell.font=Font(name='맑은 고딕',size=10)
        ws.row_dimensions[row[0].row].height=31
    if rows and keys:
        table=Table(displayName=f'Table{len(wb.worksheets)}',ref=ws.dimensions)
        table.tableStyleInfo=TableStyleInfo(name='TableStyleMedium2',showRowStripes=True)
        ws.add_table(table)
    return ws


def export(database, output):
    output.mkdir(parents=True,exist_ok=True)
    con=sqlite3.connect(database.resolve().as_uri()+'?mode=ro',uri=True)
    con.row_factory=sqlite3.Row
    info=summary(con)
    inventory=fetch(con,'SELECT * FROM v_equipment_inventory ORDER BY register_no')
    wb=Workbook()
    wb.remove(wb.active)
    readme=[
      {'항목':'목적','내용':'무인시험소 자동화 MCP 구축을 위한 장비 기준정보 DB의 검토용 내보내기'},
      {'항목':'전체 등록건수','내용':info['equipment_records']},
      {'항목':'제조사 표기 그룹','내용':info['manufacturer_groups']},
      {'항목':'제조사+모델 조합','내용':info['models']},
      {'항목':'장비종류','내용':f"{info['category_count']-1}개 종류 + 미분류 ({info['unclassified_records']}건)"},
      {'항목':'원본파일','내용':info['source']['original_filename']},
      {'항목':'원본 SHA-256','내용':info['source']['sha256']},
      {'항목':'내보내기 일시(KST)','내용':info['generated_at_kst']},
      {'항목':'전체장비','내용':'전체 등록건을 제조사·모델·장비종류별 필터로 조회. 장비ID로 다른 시트와 연결.'},
      {'항목':'제조사별칭','내용':'원문→표준표기 변환 근거. 오탈자 의심 후보는 병합하지 않고 별도 표시.'},
      {'항목':'모델별MCP조사','내용':'같은 모델을 묶어 매뉴얼·SDK·드라이버 조사에 사용. 미확인은 미지원과 다름.'},
      {'항목':'중복후보','내용':'동일 장비 확정이 아님. 일치기준별 그룹이 겹치므로 합산 금지.'},
      {'항목':'원본행','내용':'원본 9개 항목을 값 변경 없이 보존. 원본시트·원본행으로 대조.'},
      {'항목':'수정 반영','내용':'이 엑셀은 DB의 검토용 스냅샷. 편집 내용이 DB에 자동 반영되지 않음.'},
      *[{'항목':'해석 주의','내용':c} for c in info['caveats']],
    ]
    ws=write_sheet(wb,'안내',readme)
    ws.column_dimensions['A'].width=25
    ws.column_dimensions['B'].width=115
    main_keys=['equipment_id','register_no','name','manufacturer','manufacturer_raw','model_name','serial_number',
               'domain','category','source_status_hint','specification','calibration_specification','firmware','note',
               'classification_status','automation_research_hint','automation_status','lifecycle_status',
               'location','owner','calibration_due_date','manufacturer_id','model_id','sheet_name','source_row']
    write_sheet(wb,'전체장비',inventory,main_keys)
    queries={
      '제조사목록':'SELECT * FROM v_manufacturer_summary ORDER BY equipment_records DESC,manufacturer',
      '제조사별칭':'''SELECT a.raw_label,m.canonical_name,a.normalization_method,a.review_required,
           a.suggested_canonical_name,a.rationale,a.manufacturer_id FROM manufacturer_aliases a
           LEFT JOIN manufacturers m USING(manufacturer_id) ORDER BY a.review_required DESC,m.canonical_name,a.raw_label''',
      '모델별MCP조사':'SELECT * FROM v_model_research_queue ORDER BY equipment_records DESC,manufacturer,model_name',
      '장비분류':'''SELECT c.*,count(e.equipment_id) equipment_records FROM categories c
           LEFT JOIN equipment e USING(category_id) GROUP BY c.category_id ORDER BY c.domain_ko,c.category_ko''',
      '확인필요':'''SELECT q.*,e.register_no,e.name,m.canonical_name,e.model_name,e.serial_number
           FROM data_quality_issues q JOIN equipment e USING(equipment_id)
           LEFT JOIN manufacturers m USING(manufacturer_id) ORDER BY q.severity,q.equipment_id,q.issue_id''',
      '중복후보':'''SELECT g.duplicate_group_id,g.match_type,g.member_count,g.resolution_status,
           e.equipment_id,e.register_no,e.name,m.canonical_name,e.model_name,e.serial_number
           FROM duplicate_groups g JOIN duplicate_members dm USING(duplicate_group_id)
           JOIN equipment e USING(equipment_id) LEFT JOIN manufacturers m USING(manufacturer_id)
           ORDER BY g.match_type,g.duplicate_group_id,e.register_no''',
      '과거등록번호':'''SELECT i.*,e.register_no,e.name FROM equipment_identifiers i
           JOIN equipment e USING(equipment_id) ORDER BY e.register_no''',
    }
    for title,sql in queries.items():
        write_sheet(wb,title,fetch(con,sql))
    source_rows=fetch(con,'''SELECT e.equipment_id,r.sheet_name,r.row_number source_row,r.cells_json
                        FROM source_rows r JOIN equipment e USING(source_row_id) ORDER BY e.register_no''')
    raw_keys=['원본 No.','설비명 원문','제조사 원문','모델명 원문','일련번호 원문',
              '장비사양 원문','교정사양 원문','비고 원문','펌웨어 원문']
    raw=[{'equipment_id':r['equipment_id'],'sheet_name':r['sheet_name'],'source_row':r['source_row'],
          **dict(zip(raw_keys,json.loads(r['cells_json'])))} for r in source_rows]
    write_sheet(wb,'원본행',raw)
    workbook_path=output/'equipment_inventory.xlsx'
    wb.save(workbook_path)
    for filename,rows in [
        ('equipment.jsonl',inventory),
        ('manufacturers.jsonl',fetch(con,'SELECT * FROM v_manufacturer_summary ORDER BY manufacturer_id')),
        ('models.jsonl',fetch(con,'SELECT * FROM v_model_research_queue ORDER BY model_id')),
    ]:
        with (output/filename).open('w',encoding='utf-8') as stream:
            for row in rows:
                stream.write(json.dumps(row,ensure_ascii=False)+'\n')
    (output/'summary.json').write_text(json.dumps(info,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    con.close()
    print(json.dumps({'xlsx':str(workbook_path),'records':len(inventory),'sheets':wb.sheetnames},ensure_ascii=False))


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--database',type=Path,default=ROOT/'data/equipment.sqlite3')
    parser.add_argument('--output',type=Path,default=ROOT/'exports')
    args=parser.parse_args()
    export(args.database,args.output)
