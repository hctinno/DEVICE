#!/usr/bin/env python3
"""Independent source-to-database and export reconciliation. No device access."""
import argparse
from collections import Counter
import hashlib
import json
from pathlib import Path
import sqlite3

import openpyxl

ROOT=Path(__file__).resolve().parents[1]


def validate(source, database, exports):
    checks=[]

    def require(condition,message):
        if not condition:
            raise AssertionError(message)
        checks.append(message)

    con=sqlite3.connect(database.resolve().as_uri()+'?mode=ro',uri=True)
    con.row_factory=sqlite3.Row
    expected={}
    workbook=openpyxl.load_workbook(source,data_only=False)
    for sheet in workbook:
        for row in sheet.iter_rows(min_row=3):
            values=[cell.value for cell in row]
            if any(value not in (None,'') for value in values):
                expected[(sheet.title,row[0].row)]=values
    actual={(r['sheet_name'],r['row_number']):json.loads(r['cells_json'])
            for r in con.execute('SELECT * FROM source_rows')}
    require(expected==actual,'원본의 모든 데이터 셀 값·자료형과 DB 원문이 일치')
    stored_hash=con.execute('SELECT sha256 FROM source_documents').fetchone()[0]
    require(hashlib.sha256(source.read_bytes()).hexdigest()==stored_hash,'원본 파일 SHA-256 일치')
    n=len(expected)
    require(con.execute('SELECT count(*) FROM equipment').fetchone()[0]==n,'원본행과 장비 등록건수 일치')
    require(con.execute('SELECT count(*) FROM v_equipment_inventory').fetchone()[0]==n,'장비 조회 뷰에 행 증식·누락 없음')
    require(con.execute('PRAGMA integrity_check').fetchone()[0]=='ok','SQLite integrity_check 통과')
    require(not con.execute('PRAGMA foreign_key_check').fetchall(),'모든 외래키 관계 유효')
    require(not con.execute('''SELECT 1 FROM equipment e JOIN source_rows r USING(source_row_id)
              WHERE e.register_no != json_extract(r.cells_json,'$[0]')''').fetchall(),'원본 No. 연결 일치')
    require(con.execute('SELECT count(*) FROM models').fetchone()[0] ==
            con.execute('SELECT count(*) FROM model_automation_profiles').fetchone()[0],
            '모든 모델에 MCP 조사 프로필 존재')
    require(not con.execute('''SELECT 1 FROM equipment e JOIN models m USING(model_id)
                              WHERE e.manufacturer_id != m.manufacturer_id''').fetchall(),
            '모델과 장비의 제조사 연결 일치')
    require(not con.execute('SELECT 1 FROM equipment_connections').fetchall() and
            not con.execute('SELECT 1 FROM model_interfaces').fetchall() and
            not con.execute('SELECT 1 FROM model_automation_profiles WHERE remote_control_supported IS NOT NULL').fetchall(),
            '미검증 제어 인터페이스·접속주소·원격제어 지원을 임의 입력하지 않음')
    require(not con.execute("SELECT 1 FROM equipment WHERE automation_status!='not_assessed'").fetchall(),
            '모든 장비의 자동화 상태가 미확인으로 시작')
    # Independent grouping from stored source values; ensures duplicate candidates
    # preserve every original member, even when no serial is available.
    duplicate_counts=Counter(json.dumps(v[1:],ensure_ascii=False) for v in expected.values())
    exact_expected=sorted(v for v in duplicate_counts.values() if v>1)
    exact_actual=sorted(r[0] for r in con.execute("SELECT member_count FROM duplicate_groups WHERE match_type='exact_source_fields'"))
    require(exact_expected==exact_actual,'원문 8개 항목 완전일치 중복후보 그룹 재계산 일치')
    require(not con.execute('''SELECT 1 FROM duplicate_groups g JOIN duplicate_members m USING(duplicate_group_id)
                              GROUP BY g.duplicate_group_id HAVING count(*) != g.member_count''').fetchall(),
            '모든 중복후보 그룹의 구성원 수 일치')

    exported=openpyxl.load_workbook(exports/'equipment_inventory.xlsx',data_only=False,read_only=True)
    require(exported['전체장비'].max_row-1==n and exported['원본행'].max_row-1==n,
            '엑셀 전체장비·원본행 시트의 등록건수 일치')
    raw_values={}
    for row in exported['원본행'].iter_rows(min_row=2,values_only=True):
        raw_values[(row[1],row[2])]=[None if v=='' else v for v in row[3:]]
    blank_equivalent={k:[None if v=='' else v for v in values] for k,values in expected.items()}
    require(raw_values==blank_equivalent,'검토용 엑셀 원문 셀 대조 통과(빈 문자열·빈 셀 동등 취급)')
    formula_cells=[f'{s.title}!{c.coordinate}' for s in exported for row in s for c in row if c.data_type=='f']
    require(not formula_cells,'내보낸 엑셀에서 원본 문자열이 수식으로 실행되지 않음')
    with (exports/'equipment.jsonl').open(encoding='utf-8') as stream:
        records=[json.loads(line) for line in stream]
    db_inventory=[dict(r) for r in con.execute('SELECT * FROM v_equipment_inventory ORDER BY register_no')]
    require(records==db_inventory,'전체 장비 JSONL과 DB 조회 결과가 모든 필드에서 일치')
    report={'result':'passed','source_records':n,'checks_passed':len(checks),'checks':checks,
            'source_sha256':stored_hash}
    (exports/'validation.json').write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    exported.close()
    workbook.close()
    con.close()
    return report


if __name__=='__main__':
    parser=argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'data/source/equipment_register.xlsx')
    parser.add_argument('--database',type=Path,default=ROOT/'data/equipment.sqlite3')
    parser.add_argument('--exports',type=Path,default=ROOT/'exports')
    args=parser.parse_args()
    print(json.dumps(validate(args.source,args.database,args.exports),ensure_ascii=False,indent=2))
