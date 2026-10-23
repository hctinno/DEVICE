#!/usr/bin/env python3
"""Create a new inventory snapshot from an equipment register; never overwrite.

Usage: python registry/build.py --source data/source/equipment_register.xlsx
"""
import argparse
from collections import Counter, defaultdict
from datetime import datetime, timezone
import hashlib
import json
import os
from pathlib import Path
import re
import sqlite3
import tempfile

import openpyxl

from rules import CATEGORIES, brand_key, classify, clean, manufacturer_info, optional

ROOT = Path(__file__).resolve().parents[1]
HEADERS = ['No.', '설비명\nName', '제조사\nManufacturer', '모델명\nModel',
           '일련번호\nSer. No.', '장비사양\nSpecfication',
           '교정사양\nCalibration specification', '비고\nNote', '비고(펌웨어) Firmware']


def js(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':'))


def ident(prefix, value):
    return prefix + '-' + hashlib.sha256(js(value).encode()).hexdigest()[:16]


def load_source(path):
    workbook = openpyxl.load_workbook(path, data_only=False)
    records, sheets = [], []
    for sheet in workbook:
        if not any(c.value is not None for row in sheet for c in row):
            continue
        headers = [sheet.cell(2, i).value for i in range(1, 10)]
        if headers != HEADERS or sheet.max_column != 9:
            raise ValueError(f'{sheet.title}: unexpected headers or column count; nothing imported')
        sheets.append((sheet.title, 2, headers, sheet.cell(1,1).value, sheet.sheet_state))
        for row in sheet.iter_rows(min_row=3):
            cells = [c.value for c in row]
            if not any(v not in (None, '') for v in cells):
                continue
            if type(cells[0]) is not int:
                raise ValueError(f'{sheet.title}!A{row[0].row}: invalid register number')
            if any(c.data_type == 'f' for c in row):
                raise ValueError('Formula detected; reviewed formula ingestion required')
            if any(v is not None and not isinstance(v, (str,int,float,bool)) for v in cells):
                raise ValueError('Unexpected source value type; review before importing')
            records.append((sheet.title, row[0].row, cells))
    if not records:
        raise ValueError('No equipment records found')
    if len({(s, v[0]) for s, r, v in records}) != len(records):
        raise ValueError('Repeated register No. within a sheet; explicit reconciliation required')
    return records, sheets


def build(source, output, original_filename):
    if output.exists():
        raise FileExistsError(f'{output} already exists; choose a NEW snapshot path. Existing edits are protected.')
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    source_id = ident('SRC', source_hash)
    rows, sheets = load_source(source)
    output.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.tmp', delete=False) as tmp:
        temporary = Path(tmp.name)
    con = sqlite3.connect(temporary)
    con.row_factory = sqlite3.Row
    try:
        con.executescript((ROOT / 'registry/schema.sql').read_text())
        with con:
            con.execute('INSERT INTO source_documents VALUES (?,?,?,?,?,?,?)', (
                source_id, original_filename, source_hash, datetime.now(timezone.utc).isoformat(),
                str(source.resolve().relative_to(ROOT)) if source.resolve().is_relative_to(ROOT) else str(source.resolve()),
                len(sheets), len(rows)))
            for sheet, header_row, headers, title, visibility in sheets:
                con.execute('INSERT INTO source_sheets VALUES (?,?,?,?,?,?)',
                            (source_id,sheet,header_row,js(headers),title,visibility))
            con.executemany('INSERT INTO categories VALUES (?,?,?,?,?)', CATEGORIES)

            aliases = {}
            manufacturers = {}
            # Most frequent source spelling supplies the display label for generic groups.
            for raw, _ in Counter(str(v[2] or '') for _,_,v in rows).most_common():
                canonical, method, review, suggestion, rationale = manufacturer_info(raw)
                manufacturer_id = None
                if canonical:
                    key = brand_key(canonical)
                    if key not in manufacturers:
                        manufacturer_id = ident('MFR', key)
                        manufacturers[key] = manufacturer_id
                        con.execute('INSERT INTO manufacturers(manufacturer_id,canonical_name,name_key) VALUES (?,?,?)',
                                    (manufacturer_id,canonical,key))
                    manufacturer_id = manufacturers[key]
                aliases[raw] = (manufacturer_id, method, review)
                con.execute('INSERT INTO manufacturer_aliases VALUES (?,?,?,?,?,?)',
                            (raw,manufacturer_id,method,int(review),suggestion,rationale))

            model_ids = {}
            duplicate_keys = {kind:defaultdict(list) for kind in ('exact_source_fields','manufacturer_model_serial','serial_only')}
            for ordinal, (sheet, row_number, cells) in enumerate(rows,1):
                row_id = ident('ROW', [source_hash,sheet,row_number])
                con.execute('INSERT INTO source_rows VALUES (?,?,?,?,?)',
                            (row_id,source_id,sheet,row_number,js(cells)))
                equipment_id = f'EQ-{ordinal:06d}'
                register_no = cells[0]
                name, _, model, serial, spec, cal_spec, note, firmware = [optional(x) for x in cells[1:]]
                manufacturer_id, normalization_method, manufacturer_review = aliases[str(cells[2] or '')]
                model_id = None
                if manufacturer_id and model:
                    key = (manufacturer_id, model.casefold())
                    if key not in model_ids:
                        model_ids[key] = ident('MOD', key)
                        con.execute('INSERT INTO models VALUES (?,?,?,?)',
                                    (model_ids[key],manufacturer_id,model,model.casefold()))
                        con.execute('INSERT INTO model_automation_profiles(model_id) VALUES (?)', (model_ids[key],))
                    model_id = model_ids[key]
                category_id, classification_status, rule = classify(name)
                status_hint = 'no_status_evidence'
                if note and re.search(r'폐기|사용\s*중지|고장|수리|불량|분실',note):
                    status_hint = 'status_history_review'
                    if re.fullmatch(r'폐기(?:\([^)]*\))?[.!]?', note):
                        status_hint = 'retired_reported'
                    elif re.fullmatch(r'사용\s*중지(?:\([^)]*\))?[.!]?', note):
                        status_hint = 'out_of_service_reported'
                con.execute('''INSERT INTO equipment(
                    equipment_id,source_row_id,register_no,name,manufacturer_id,model_id,
                    model_name,serial_number,specification,calibration_specification,note,firmware,
                    category_id,classification_status,classification_rule,source_status_hint)
                    VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)''', (
                    equipment_id,row_id,register_no,name,manufacturer_id,model_id,model,serial,
                    spec,cal_spec,note,firmware,category_id,classification_status,rule,status_hint))

                def issue(field, kind, severity, detail):
                    con.execute('''INSERT INTO data_quality_issues
                      (equipment_id,field_name,issue_type,severity,detail) VALUES (?,?,?,?,?)''',
                                (equipment_id,field,kind,severity,detail))
                for field,value,severity in [('name',name,'high'),('manufacturer',manufacturer_id,'high'),
                                             ('model',model,'medium'),('serial_number',serial,'medium')]:
                    if value is None:
                        issue(field,'missing_identifier',severity,'원본 공란/결측/미지정; 자동 보완하지 않음')
                if manufacturer_review and manufacturer_id:
                    issue('manufacturer','manufacturer_alias_review','medium',
                          '제조사 오탈자 의심; 원문 그룹을 유지하고 별칭 후보만 제안')
                if normalization_method in ('non_manufacturer','multiple_manufacturers'):
                    issue('manufacturer',normalization_method,'medium','제조사 별칭표의 원문과 사유 확인 필요')
                if classification_status == 'unclassified':
                    issue('category','unclassified','low','이름 규칙만으로 분류 불가; 용도 확인 필요')
                if status_hint != 'no_status_evidence':
                    issue('lifecycle_status','lifecycle_note_review','medium',
                          '폐기·사용중지·고장·수리 등 비고 이력 존재; 현재 운용 상태를 현장 확인해야 함')
                if model is None or manufacturer_id is None:
                    issue('model_id','model_catalog_unresolved','medium',
                          '제조사와 모델을 함께 특정할 수 없어 모델 마스터에 연결하지 않음')

                if note:
                    for match in re.finditer(r'구\s*등록번호\s*[:：]?\s*([A-Za-z0-9][A-Za-z0-9._/-]*)', note):
                        con.execute('''INSERT INTO equipment_identifiers
                          (equipment_id,identifier_type,identifier_value,source_field,source_excerpt)
                          VALUES (?,?,?,?,?)''', (equipment_id,'legacy_registration',match.group(1),'note',match.group(0)))

                duplicate_keys['exact_source_fields'][js(cells[1:])].append(equipment_id)
                if manufacturer_id and model and serial:
                    duplicate_keys['manufacturer_model_serial'][js([manufacturer_id,model.casefold(),serial.casefold()])].append(equipment_id)
                if serial:
                    duplicate_keys['serial_only'][js(serial.casefold())].append(equipment_id)

            for kind, groups in duplicate_keys.items():
                for key, members in groups.items():
                    if len(members) < 2:
                        continue
                    group_id = ident('DUP', [kind,key])
                    con.execute('INSERT INTO duplicate_groups(duplicate_group_id,match_type,match_key_json,member_count) VALUES (?,?,?,?)',
                                (group_id,kind,key,len(members)))
                    con.executemany('INSERT INTO duplicate_members VALUES (?,?)', [(group_id,e) for e in members])
            assert con.execute('PRAGMA integrity_check').fetchone()[0] == 'ok'
            assert not con.execute('PRAGMA foreign_key_check').fetchall()
            assert con.execute('SELECT COUNT(*) FROM equipment').fetchone()[0] == len(rows)
        con.close()
        # Hard-link publication is atomic and fails if another builder created output.
        os.link(temporary, output)
        temporary.unlink()
        return {'records':len(rows),'manufacturers':len(manufacturers),'models':len(model_ids),
                'sha256':source_hash,'database':str(output)}
    except BaseException:
        con.close()
        temporary.unlink(missing_ok=True)
        raise


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--source',type=Path,default=ROOT/'data/source/equipment_register.xlsx')
    parser.add_argument('--output',type=Path,default=ROOT/'data/equipment.sqlite3')
    parser.add_argument('--original-filename',default='시험설비_관리대장 (16).xlsx')
    args=parser.parse_args()
    print(json.dumps(build(args.source,args.output,args.original_filename),ensure_ascii=False,indent=2))


if __name__ == '__main__':
    main()
