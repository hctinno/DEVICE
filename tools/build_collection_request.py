"""Build an external-agent collection request from the preserved source inventories."""
from pathlib import Path
import csv
import hashlib
import json
import re
import sqlite3
from urllib.parse import quote_plus
import zipfile

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'handoff'


def search(query):
    return 'https://www.google.com/search?q=' + quote_plus(query)


def write_csv(name, rows, headers=None):
    headers = headers or list(rows[0])
    with (OUT / name).open('w', encoding='utf-8-sig', newline='') as stream:
        writer = csv.DictWriter(stream, fieldnames=headers)
        writer.writeheader()
        writer.writerows(rows)


def build():
    OUT.mkdir(exist_ok=True)
    requests = []

    def add(identifier, group, title, priority, url, query, scope, instruction,
            link_type='official_page', verified='official_route_not_live_checked', edition='resolve_current_and_applicable'):
        requests.append(dict(request_id=identifier, group=group, title=title,
            priority=priority, edition_to_collect=edition, official_url=url,
            direct_download_url=url if link_type == 'direct_file' else '',
            discovery_url=search(query), link_type=link_type,
            link_verification=verified, scope=scope, collection_instructions=instruction,
            current_acquisition='original_not_acquired', owner='external_collection_agent'))

    for part in [1, 2, 5, 11, 15, 18, 20, 22, 24, 25, 27, 30, 73, 74, 80, 87, 90, 95, 96, 97, 101]:
        add(f'FCC-PART-{part}', 'FCC regulations', f'47 CFR Part {part}',
            'P0' if part in [1, 2, 15, 18] else 'P1',
            f'https://www.ecfr.gov/current/title-47/part-{part}',
            f'site:ecfr.gov 47 CFR Part {part}', 'FCC scope inventory',
            'Collect complete current XML/HTML and readable PDF if offered, amendments, effective dates, transitions, reserved/removed clauses and IBR references. Resolve Federal Register amendments; do not substitute an old annual CFR without labeling.')
    for section in ['1.1307','1.1310','1.50002','2.803','2.903','2.906','2.907','2.909','2.910','2.911','2.924','2.925','2.926','2.929','2.933','2.938','2.943','2.944','2.945','2.947','2.948','2.949','2.950','2.960','2.962','2.1033','2.1041','2.1043','2.1046','2.1047','2.1049','2.1051','2.1053','2.1055','2.1057','2.1077','2.1091','2.1093','15.31','15.33','15.35','15.37','15.38','15.101','15.107','15.109','15.201','15.203','15.204','15.205','15.207','15.209','15.212','15.225','15.231','15.236','15.247','15.249','15.255','15.256','15.407']:
        add('FCC-SEC-'+section, 'FCC clause and references', '47 CFR '+section, 'P0',
            'https://www.ecfr.gov/current/title-47/section-'+section,
            'site:ecfr.gov "'+section+'"', 'FCC common / selected RF and EMC',
            'Index this clause inside its complete Part snapshot; collect referenced sources and transition dates. Shared source files may satisfy multiple request_ids. Applicability remains product-specific.')

    text=(ROOT/'data/regulatory/fcc/extracts/G03_search_index.txt').read_text()
    kdb={}
    for number, question in re.findall(r'^([0-9]{6})\s+(.+)$',text,re.M):
        kdb.setdefault(number,question)
    kdb.setdefault('484596','Referencing test data from another equipment authorization application')
    critical=set('558074 789033 987594 905462 447498 865664 248227 996369 974614 662911 594280 388624 484596 935210 940660 653005 680106 227764 171051 414788 822428 746324 300643 178919 784748 896810 941225 442401 641163 853844 986446'.split())
    known={'558074':'21124','300643':'21079','996369':'44637'}
    for number,question in sorted(kdb.items()):
        url=('https://apps.fcc.gov/oetcf/kdb/forms/FTSSearchResultPage.cfm?id='+known[number]+'&switch=P') if number in known else 'https://apps.fcc.gov/oetcf/kdb/index.cfm'
        add('KDB-'+number,'KDB publication and ALL attachments',number+' '+question,
            'P0' if number in critical else 'P1',url,
            f'site:apps.fcc.gov/oetcf/kdb "{number}"','FCC guidance discovery seed',
            'Search this publication number in the official KDB. Save publication page and every applicable D01/D02/etc attachment, appendices, spreadsheets, change notes and replacement links. Separate current, draft, superseded and withdrawn. Seed came from an old index; verify current status and recursively discover new publications.',
            link_type='official_page' if number in known else 'official_portal',
            verified='official_search_index_seen_current_attachments_unverified')

    standards=[
        ('ANSI C63.10','2020','https://standards.ieee.org/ieee/C63.10/10445/','15.31/15.37/15.38; intentional radiators'),
        ('ANSI C63.10 Cor1','2023','https://standards.ieee.org/ieee/C63.10-2020_Cor1/10915/','Correction to 2020; verify FCC treatment'),
        ('ANSI C63.10a','2024','https://standards.ieee.org/ieee/C63.10a/11614/','Amendment; do not auto-substitute for incorporated edition'),
        ('ANSI C63.4','2014','https://standards.ieee.org/ieee/C63.4/5841/','15.31 exclusions and 2.948 site clauses'),
        ('ANSI C63.4a','2017','','2.910 / 2.948 site validation'),
        ('ANSI C63.25.1','2018','','2.910 / 2.948 1-18 GHz site validation'),
        ('ANSI C63.26','2015','','2.910 / licensed transmitters'),
        ('ANSI C63.17','2013','','15.38 / UPCS'),
        ('CISPR 16-1-4','2010-04 Edition 3.0','','2.910 / 2.948 site reference'),
        ('ISO/IEC 17025','2017','https://www.iso.org/standard/66912.html','Laboratory competence / accreditation'),
        ('ISO/IEC 17025','2005','https://www.iso.org/standard/39883.html','Historical incorporated edition; retrieve transition evidence'),
        ('ISO/IEC 17011','2004 and current referenced edition','','Accreditation body / 2.910; distinguish updated editions'),
        ('ISO/IEC 17065','2012','','TCB/product certification body / 2.910'),
        ('ANSI C63.19','edition required by applicable FCC HAC rules','','Hearing aid compatibility, only where applicable'),
        ('IEEE 1528','edition accepted by current SAR guidance','','SAR measurement reference; verify current transition'),
        ('IEC/IEEE 62209-1528','edition accepted by current FCC guidance','','SAR methodology candidate; verify acceptance'),
        ('IEC 62209-1 and IEC 62209-2','historical/current references as applicable','','SAR historical replay and reference chains'),
        ('IEC/IEEE 63195-1 and 63195-2','edition accepted by current FCC guidance','','Power density measurement/computation candidate'),
        ('ANSI C63.23','edition referenced by applicable measurement method','','Uncertainty / measurement method reference chain'),
        ('CISPR 16-1-1','edition referenced by applicable method','','Receiver/detector instrument requirements'),
        ('CISPR 16-2-1 and 16-2-3','edition referenced by applicable method','','Conducted/radiated method reference chain'),
        ('CISPR 16-4-2','edition referenced by applicable method','','Measurement uncertainty reference chain'),
        ('CISPR 22','1997 Third Edition','','15.38 and 15.109 conditional historical alternative'),
        ('ETSI EN 300 422-1','V2.2.1 (2021-11)','https://www.etsi.org/deliver/etsi_en/300400_300499/30042201/02.02.01_60/en_30042201v020201p.pdf','15.38 / wireless microphones; validate applicability'),
    ]
    for n,(title,edition,url,scope) in enumerate(standards,1):
        domain='standards.ieee.org' if ('ANSI' in title or title.startswith('IEEE')) else 'webstore.iec.ch' if ('IEC' in title or 'CISPR' in title) else 'iso.org'
        if title.startswith('ISO'):domain='iso.org'
        add(f'STD-{n:03d}','Measurement / quality standards',title,'P0' if n<=13 else 'P1',url or ('https://'+domain+'/'),
            f'site:{domain} "{title}" {edition}',scope,
            'Obtain the named applicable edition and all corrigenda/errata plus catalog and access terms. Newest publisher edition and legally applicable edition are separate. Use authorized public/owned license access; return purchase/login/read-only status when original bytes cannot be provided. Follow normative references recursively. Candidate standards are not automatically FCC accepted.',
            link_type='direct_file' if url.endswith('.pdf') else 'official_page' if url else 'official_portal',
            verified='official_reference_or_catalog_seen' if n<=13 or url else 'candidate_verify_applicability',edition=edition)

    specials=[
        ('DISC-KDB','Current KDB full catalog, major guidance, recent updates and draft list','https://apps.fcc.gov/oetcf/kdb/reports/GuidedPublicationList.cfm','site:apps.fcc.gov KDB major guidance','Enumerate the live catalog and record added/removed/replaced IDs compared with seed; follow all applicable references. Return completeness report; 179 seed entries are not a current exhaustive count.'),
        ('DISC-FR','FCC Orders, Public Notices, Federal Register amendments and transition notices','https://www.federalregister.gov/agencies/federal-communications-commission','site:federalregister.gov FCC equipment authorization measurement standards','Follow citations from every requested CFR/KDB; save final and correction/effective-date notices. Distinguish NPRM/draft/proposed from effective rules.'),
        ('DISC-EA','Equipment authorization procedures, forms, exhibit/label/confidentiality guidance','https://www.fcc.gov/engineering-technology/laboratory-division/general/equipment-authorization','site:fcc.gov equipment authorization procedures exhibits form 731','Collect instructions, current form/exhibit categories and SDoC/Certification distinctions.'),
        ('DISC-WORKSHOP','FCC/TCB workshop presentations and measurement updates','https://www.fcc.gov/','site:fcc.gov equipment authorization presentations TCB workshop','Find official presentation archive; collect recent and transition-relevant slides and cited documents. Slides are explanatory guidance, not law.'),
        ('DISC-LABS','FCC recognized testing laboratories and accreditation bodies; scope procedures','https://apps.fcc.gov/oetcf/eas/reports/TestFirmSearch.cfm','site:fcc.gov testing laboratory qualifications recognized accreditation','Save current recognition/scope procedures and lookup instructions; do not infer this laboratory is recognized.'),
        ('DISC-MP5','FCC MP-5 / Part 18 ISM measurement procedures and replacements','https://www.fcc.gov/','site:fcc.gov "MP-5" measurements','Locate official measurement publication, history and any replacements; verify current acceptance per Part 18.'),
        ('DISC-EXPOSURE','FCC OET Bulletin 65 and supplements / current RF exposure replacements','https://www.fcc.gov/','site:fcc.gov OET Bulletin 65 RF exposure','Preserve historical labels and supersession evidence; current KDB and CFR determine applicable method.'),
        ('DISC-IBR','IEEE reading room / ANSI IBR / publisher licensing routes','https://standards.ieee.org/access-standards/','site:ibr.ansi.org standards read only','Return per-standard access conditions; read-only screen capture is not an original reusable standard.'),
        ('DISC-REPORTS','Public FCC filing examples with RF, EMC, SAR/PD, DFS and 6 GHz exhibits','https://apps.fcc.gov/oetcf/eas/reports/GenericSearch.cfm','site:apps.fcc.gov equipment authorization test report','Collect representative official filing exhibits with FCC ID/date/technology. Label examples, not normative rules or original instrument data; record confidentiality exclusions. Cover Part15B, DTS, FHSS, U-NII/DFS/6GHz, modules, Part18, licensed and exposure.'),
        ('DEV-MCP','MCP protocol specification and official Python SDK','https://modelcontextprotocol.io/specification/latest','site:modelcontextprotocol.io specification tools','Save versioned specification, SDK examples, version and license. Pin protocol version, not a moving latest label.'),
        ('DEV-VISA','VISA/PyVISA backend, vendor driver and SCPI transport documentation','https://pyvisa.readthedocs.io/en/latest/','PyVISA VISA IVI HiSLIP VXI-11 USBTMC official specification','Collect API/transport docs and licenses; identify 32/64-bit OS compatibility. Store installers only where permitted; do not run them.'),
        ('DEV-IVI','IVI Foundation specifications, SCPI and IVI driver interfaces','https://www.ivifoundation.org/','site:ivifoundation.org specifications HiSLIP IVI','Collect referenced transport/driver specs and version history, without assuming equipment supports them.'),
    ]
    for ident,title,url,query,instruction in specials:
        add(ident,'Discovery / supporting documents',title,'P0',url,query,'All pipeline stages',instruction,link_type='official_portal')
    future=[
        ('KC-REG','KC/RRA conformity assessment notices, technical regulations, test methods, forms','https://www.rra.go.kr/','site:rra.go.kr 적합성평가 고시 시험방법','KC: resolve law.go.kr/RRA notices and KS editions; no automatic reuse of FCC limits'),
        ('KC-KS','Referenced KS/IEC/CISPR standards and permitted original access','https://standard.go.kr/','site:standard.go.kr KS 전자파 시험','KC standards and corrigenda; license-aware access'),
        ('EU-RED','RED 2014/53/EU consolidated text, delegated acts and current harmonized-standard decisions','https://eur-lex.europa.eu/eli/dir/2014/53/oj','site:eur-lex.europa.eu RED harmonised standards 2014/53','EU radio; gather OJ restrictions and cessation dates, not just latest ETSI editions'),
        ('EU-EMC','EMC 2014/30/EU and current harmonized-standard decisions','https://eur-lex.europa.eu/eli/dir/2014/30/oj','site:eur-lex.europa.eu 2014/30 harmonised standards','EU EMC; emission and immunity distinction'),
        ('EU-ETSI','EN300328, EN301893, EN303687, EN301489-1/-17 and applicable technology families','https://www.etsi.org/standards','site:etsi.org deliver EN 300 328 301 893 303 687 301 489','Collect OJ applicable plus newest versions and change history; product-dependent family expansion'),
        ('EU-EMC-STD','EN55032/55035, applicable product-family standards and IEC61000-4 series','https://webstore.iec.ch/','IEC CISPR 32 CISPR 35 IEC 61000-4 official','Future EU references; FCC general pipeline must not assume immunity tests are always FCC requirements'),
    ]
    for ident,title,url,query,instruction in future:
        add(ident,'Future market expansion',title,'P2',url,query,'After FCC core',instruction,link_type='official_portal')
    write_csv('regulatory_requests.csv',requests)

    db=sqlite3.connect((ROOT/'data/equipment.sqlite3').as_uri()+'?mode=ro',uri=True);db.row_factory=sqlite3.Row
    portal_map={
        'Rohde & Schwarz':('https://www.rohde-schwarz.com/manual/','rohde-schwarz.com'),
        'Keysight':('https://www.keysight.com/us/en/support.html','keysight.com'),
        'Agilent':('https://www.keysight.com/us/en/support.html','keysight.com'),
        'HP':('https://www.keysight.com/us/en/support.html','keysight.com'),
        'SPEAG':('https://speag.swiss/downloads','speag.swiss'),
        'SCHWARZBECK':('https://www.schwarzbeck.de/','schwarzbeck.de'),
        'Tektronix':('https://www.tek.com/en/support/datasheets-manuals-software-downloads','tek.com'),
        'Anritsu':('https://www.anritsu.com/','anritsu.com'),
        'ETS-Lindgren':('https://www.ets-lindgren.com/','ets-lindgren.com'),
        'YOKOGAWA':('https://tmi.yokogawa.com/','tmi.yokogawa.com'),
        'HIOKI':('https://www.hioki.com/','hioki.com'),
        'KIKUSUI':('https://global.kikusui.co.jp/','kikusui.co.jp'),
    }
    models=[]
    sql='''SELECT m.model_id,m.model_name,f.canonical_name,f.manufacturer_id,
      GROUP_CONCAT(DISTINCT e.name) names,GROUP_CONCAT(DISTINCT c.category_ko) categories,
      COUNT(*) registration_rows FROM models m JOIN manufacturers f USING(manufacturer_id)
      JOIN equipment e USING(model_id) JOIN categories c USING(category_id) GROUP BY m.model_id ORDER BY f.canonical_name,m.model_name'''
    for r in db.execute(sql):
        manufacturer=r['canonical_name'];model=r['model_name'];portal,domain=portal_map.get(manufacturer,('',''))
        prefix=('site:'+domain+' ') if domain else ''
        models.append(dict(request_id='EQ-'+r['model_id'],model_id=r['model_id'],manufacturer_id=r['manufacturer_id'],manufacturer=manufacturer,model=model,equipment_names=r['names'],categories=r['categories'],registration_rows=r['registration_rows'],priority='P1' if re.search('신호|분석|전력|안테나|LISN|SAR|무선|증폭|챔버|위치',r['categories']) else 'P2',official_portal=portal,portal_verification='official_route_candidate_verify_identity' if portal else 'manufacturer_identity_lookup_required',manual_search_url=search(prefix+'"'+manufacturer+'" "'+model+'" manual datasheet'),programming_search_url=search(prefix+'"'+model+'" programming remote control SCPI API SDK'),document_types='datasheet; operating/user; programming/remote control if supported; communication protocol; SDK/driver and examples; firmware compatibility; calibration/verification procedure; applicable option manuals',match_requirements='Exact model/generation/options/firmware applicability; preserve legacy brand; verify manufacturer identity; passive accessories may have datasheet only. Manuals do not establish unit calibration or live remote-control support.',result_status='pending'))
    write_csv('equipment_model_requests.csv',models)
    unresolved=[dict(r) for r in db.execute('''SELECT e.equipment_id,e.register_no,e.name,m.canonical_name manufacturer,e.model_name,e.serial_number,e.specification FROM equipment e LEFT JOIN manufacturers m USING(manufacturer_id) WHERE e.model_id IS NULL ORDER BY e.register_no''')]
    for r in unresolved:r.update(request_id='IDENTIFY-'+r['equipment_id'],action='Resolve from original register/nameplate/owner; do not guess manufacturer/model or attach a similar manual',lookup_url=search(' '.join(str(r[k] or '') for k in ['manufacturer','model_name','name'])+' manual'),owner='collection_agent_then_lab_if_unresolved')
    write_csv('unresolved_equipment.csv',unresolved)
    assets=[dict(r) for r in db.execute('''SELECT e.equipment_id,e.register_no,e.model_id,e.name,m.canonical_name manufacturer,e.model_name,e.serial_number,e.source_status_hint,e.lifecycle_status FROM equipment e LEFT JOIN manufacturers m USING(manufacturer_id) ORDER BY register_no''')]
    for r in assets:r.update(request_id='ASSET-'+r['equipment_id'],required_evidence='Confirm active unit and scope; actual options/firmware/location; applicable calibration certificate/correction data/verification; connection identity if available',acquisition_route='lab_owned_not_public_download',note='Registration row is not a verified distinct working unit; resolve duplicates and lifecycle before requiring calibration')
    write_csv('asset_evidence_requests.csv',assets)
    db.close()
    known_links=[
        ('Rohde & Schwarz','NRP6A','user manual family','https://scdn.rohde-schwarz.com/ur/pws/dl_downloads/pdm/cl_manuals/user_manual/1177_6017_01/NRPxxAN_UserManual_en_17.pdf'),
        ('Agilent / Keysight','E3632A','user guide','https://www.keysight.com/us/en/assets/9018-01309/user-manuals/9018-01309.pdf'),
        ('Agilent / Keysight','E3632A','service guide','https://www.keysight.com/us/en/assets/9018-01459/user-manuals/9018-01459.pdf'),
        ('Agilent / Keysight','N9020A','real-time option user/programmer reference','https://www.keysight.com/zz/en/assets/9018-03890/user-manuals/9018-03890.pdf'),
        ('Agilent / Keysight','N9020A','service guide','https://www.keysight.com/zz/en/assets/9018-04008/service-manuals/9018-04008.pdf'),
        ('SPEAG','EX3DV4','product page','https://speag.swiss/components/probes-2/dosimetric/ex3dv4-isotropic-dos-probe-3'),
        ('SPEAG','EX3DV4 / DASY','6-10 GHz application note','https://speag.swiss/assets/downloads/products/dasy/application-notes/Measurements-6-10-GHz/AppNote_6-10GHz_230810.pdf'),
        ('SCHWARZBECK','NNHV8123 / NNBM8124 families','family manual / exact suffix must be matched','https://www.schwarzbeck.de/Datenblatt/k8123.pdf'),
    ]
    write_csv('located_equipment_links.csv',[dict(manufacturer=a,model=b,document=c,url=d,link_type='direct_file' if d.endswith('.pdf') else 'official_page',verification='official_web_search_result_seen_2026-10-23_bytes_not_downloaded',remaining='Check title-page model list, options, revision and whether newer applicable version exists') for a,b,c,d in known_links])
    internal=[]
    for n,(title,why) in enumerate([
        ('시험소 인정·FCC recognition·scope','인정기관, 사이트, 범위, 유효성 확인'),('사이트 검증·챔버 도면·NSA/SVSWR 등','시험거리·주파수·배치별 검증 범위'),('교정·보정·중간점검·추적성','실제 장비/주파수/시점과 연결'),('방법별 불확도 예산·판정 정책','범위별 계산·보고·판정 재현'),('LIMS·시험 SW·보고 SW API/내보내기 문서','버전·권한·파일 형식·라이선스·테스트 환경'),('DUT 시험모드 API와 펌웨어·안테나·동시송신','제품별 채널/전력/모드 설정과 읽기 검증'),('시료·배선·지그·마스트·턴테이블·RF 스위치','실제 위치/경로/한계와 반복성'),('현장 제어·인터록·긴급정지·복구','AI 연결과 무관한 현장 보호/복구 검증'),('제품 자료 및 실제 시료 식별','HW/FW/모델/블록도/회로도/사용설명서/전원'),('실측 원시 데이터와 셋업·계산·보고서 연결 사례','규격 중심 새 형식의 독립 재현 검증'),('기술 검토·발행 권한·서명·개정 관리','정식 보고서 운영과 초안 생성 구별'),('백업·기록 보존·고객 접근 분리','원본 복구 및 고객별 권한 확인'),('현장 네트워크 주소·프로토콜·장비 식별 응답','검증된 연결 설정; 비밀번호/API 키는 전달 ZIP에 넣지 않음'),('운영범위·제품군·최악 조건·외주 범위','규정별 실제 시험 능력과 책임 분담'),('QMS·문서관리·숙련도/비교시험·부적합 절차','시험소 운영 체계와 지속 검증'),('기존 매뉴얼/표준 라이선스 보유 목록','중복 구매 없이 허용된 원문 접근')],1):
        internal.append(dict(request_id=f'LOCAL-{n:03d}',document=title,purpose=why,priority='P1',owner='lab_or_product_owner',download_url='',status='not_publicly_reconstructible',instruction='Collect only when available and authorized; not a prerequisite for public-source intake or software construction'))
    write_csv('lab_private_inputs.csv',internal)

    counts={'regulatory_requests':len(requests),'kdb_publications':len(kdb),'standard_requests':len(standards),'equipment_models':len(models),'unresolved_equipment_rows':len(unresolved),'asset_rows':len(assets),'private_input_groups':len(internal),'located_equipment_links':len(known_links)}
    assert len(models)==2867 and len(assets)==6470 and len(unresolved)==270
    assert len({r['request_id'] for r in requests+models+unresolved+assets+internal})==sum(map(len,[requests,models,unresolved,assets,internal]))
    (OUT/'counts.json').write_text(json.dumps(counts,ensure_ascii=False,indent=2)+'\n')
    wb=Workbook();wb.remove(wb.active)
    for name,sheet in [('regulatory_requests.csv','규정_KDB_표준'),('equipment_model_requests.csv','모델별_매뉴얼'),('located_equipment_links.csv','발견된_직접링크'),('unresolved_equipment.csv','장비식별_보완'),('asset_evidence_requests.csv','실물별_내부증거'),('lab_private_inputs.csv','시험소_내부입력')]:
        ws=wb.create_sheet(sheet)
        with (OUT/name).open(encoding='utf-8-sig',newline='') as stream:
            for row in csv.reader(stream):
                ws.append(row)
                for cell in ws[ws.max_row]:
                    # A source label is data, including strings beginning with '='.
                    cell.data_type='s'
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for i in range(1,ws.max_column+1):ws.column_dimensions[get_column_letter(i)].width=26
        for cell in ws[1]:cell.fill=PatternFill('solid',fgColor='17365D');cell.font=Font(bold=True,color='FFFFFF')
        for row in ws.iter_rows(min_row=2):
            for cell in row:cell.alignment=Alignment(vertical='top',wrap_text=True)
    wb.save(OUT/'collection_requests.xlsx')
    files=sorted(p for p in OUT.rglob('*') if p.is_file() and p.name!='bundle_manifest.json')
    manifest={str(p.relative_to(OUT)):hashlib.sha256(p.read_bytes()).hexdigest() for p in files}
    (OUT/'bundle_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+'\n')
    with zipfile.ZipFile(ROOT/'exports/agent_collection_request.zip','w',zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.rglob('*')):
            if p.is_file():z.write(p,p.relative_to(OUT))
    print(json.dumps(counts,ensure_ascii=False))


if __name__=='__main__':build()
