from pathlib import Path
import json, sqlite3, hashlib, collections, zipfile, re, datetime
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'deliverables'

def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()

def main():
    inventory = json.loads((ROOT / 'inventory_items.json').read_text())
    items = json.loads((OUT / 'model_control_sources.json').read_text())
    meta = json.loads((ROOT / 'inventory_snapshot.json').read_text())
    assert len(items) == len(inventory) == 3137
    assert collections.Counter(r['item_id'] for r in inventory) == collections.Counter(r['item_id'] for r in items)
    assert all(len(r['equipment_ids']) == len(r['register_numbers']) == r['registration_rows'] for r in items)
    assert all(r['research_status'] != 'not_searched' for r in items)
    errors = []
    source_urls = set()
    for r in items:
        if r['research_status'] == 'exact_source' and not any(s['match_level'] == 'exact_model' for s in r['sources']):
            errors.append({'item_id':r['item_id'], 'issue':'exact_status_without_exact_evidence'})
        for s in r['sources']:
            source_urls.add(s['url'])
            if '...' in s['url'] or '…' in s['url'] or re.search(r'\s', s['url']):
                errors.append({'item_id':r['item_id'], 'issue':'URL_format', 'url':s['url']})
    (ROOT / 'quality_warnings.json').write_text(json.dumps(errors, ensure_ascii=False, indent=2))
    assert not errors, f'{len(errors)} quality issues remain'
    original = Path('/workspace/DEVICE/data/equipment.sqlite3')
    assert sha(original) == meta['database_sha256'], 'original register DB changed'
    orig_db = sqlite3.connect('file:' + str(original) + '?mode=ro', uri=True)
    original_ids = {x[0] for x in orig_db.execute('select equipment_id from equipment')}
    orig_db.close()
    mapped = [x for r in items for x in r['equipment_ids']]
    assert len(mapped) == len(set(mapped)) == 6470
    assert set(mapped) == original_ids
    db = sqlite3.connect(OUT / 'equipment_control_research.sqlite3')
    assert db.execute('pragma integrity_check').fetchone()[0] == 'ok'
    assert db.execute('pragma foreign_key_check').fetchall() == []
    assert {x[0] for x in db.execute('select equipment_id from register_mapping')} == original_ids
    assert db.execute('select count(*) from models').fetchone()[0] == 3137
    db.close()
    wb = load_workbook(OUT / '전체장비_제어자료_조사.xlsx', read_only=True)
    assert wb['전체대조'].max_row - 1 == 6470
    assert sum(wb[n].max_row - 1 for n in ['제어대상', '확인필요', '제외항목']) == 3137
    formula_count = sum(c.data_type == 'f' for w in wb for row in w for c in row)
    assert formula_count == 0
    wb.close()
    checks = json.loads((ROOT / 'code_link_checks.json').read_text())
    checked = {r['source_url']:r for r in checks}
    wanted = {s['url'] for r in items for s in r['sources'] if s['kind'] == 'source_code' and re.match(r'https://(?:github\.com|raw\.githubusercontent\.com)/', s['url'])}
    assert wanted <= set(checked), 'new GitHub code URLs need HTTP checking'
    http_failures = [checked[u] for u in sorted(wanted) if checked[u].get('http_status') != 200]
    status = {
        'validated_at_utc':datetime.datetime.now(datetime.timezone.utc).isoformat(),
        'register_rows':6470, 'items':3137,
        'coverage':'all original equipment IDs represented exactly once',
        'unsearched_items':0, 'url_format_issues':0, 'exact_source_status_issues':0,
        'original_database_unchanged':True,
        'sqlite_integrity':'ok', 'sqlite_foreign_key_issues':0,
        'xlsx_all_rows_verified':True, 'xlsx_formula_cells':0,
        'github_code_urls_checked':len(wanted),
        'github_code_urls_http200':len(wanted)-len(http_failures),
        'github_code_url_failures':http_failures,
        'all_external_urls_downloaded':False,
        'hardware_tested':False,
    }
    (OUT / 'final_validation.json').write_text(json.dumps(status, ensure_ascii=False, indent=2))
    manifest = {p.name:{'bytes':p.stat().st_size,'sha256':sha(p)} for p in sorted(OUT.iterdir()) if p.is_file() and p.name != 'manifest.json'}
    (OUT / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2))
    archive = ROOT / 'equipment-control-research-2026-10-24.zip'
    with zipfile.ZipFile(archive, 'w', compression=zipfile.ZIP_DEFLATED) as z:
        for p in sorted(OUT.iterdir()):
            if p.is_file():z.write(p, p.relative_to(OUT))
        # Preserve original field-level research and web logs, without repackaging downloaded third-party code.
        paths = list(ROOT.glob('*.py')) + [ROOT/'RESEARCH_PROTOCOL.md', ROOT/'inventory_snapshot.json', ROOT/'inventory_items.json', ROOT/'code_link_checks.json']
        paths += list((ROOT/'inputs').glob('*.json'))
        for p in (ROOT/'results').rglob('*.json'):
            rel = p.relative_to(ROOT/'results')
            if any(part in ('repos','repositories','source_reads') for part in rel.parts):continue
            paths.append(p)
        paths += list((ROOT/'results').glob('*/README.md'))
        paths += list((ROOT/'results').glob('*/README.ko.md'))
        paths += list((ROOT/'results').glob('*/summary.md'))
        paths += list((ROOT/'root_review').glob('*.json'))
        for p in sorted(set(paths)):
            z.write(p, 'provenance/' + str(p.relative_to(ROOT)))
    with zipfile.ZipFile(archive) as z:
        assert z.testzip() is None
        for name in manifest:assert name in z.namelist()
    print(json.dumps(dict(status,archive=str(archive),archive_bytes=archive.stat().st_size,archive_sha256=sha(archive)),ensure_ascii=False,indent=2))

if __name__ == '__main__':
    main()
