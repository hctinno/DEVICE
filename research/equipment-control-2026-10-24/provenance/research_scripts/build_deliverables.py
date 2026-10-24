from pathlib import Path
import json,sqlite3,collections,csv,hashlib,datetime,sys,re
from openpyxl import Workbook,load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.worksheet.table import Table,TableStyleInfo
ROOT=Path(__file__).resolve().parent
OUT=ROOT/'deliverables'
SCOPE={'included':'제어·계측 대상','excluded_passive':'제어 비대상(부품·치구 등)','identity_unresolved':'제조사·모델 식별 보류','needs_scope_review':'제어 대상 여부 확인'}
STATUS={'exact_source':'모델명 일치 자료','family_source':'제품군 자료','vendor_only':'제조사/제품 참고 자료','searched_no_public_control_source':'공개 제어 자료 미확보','identity_unresolved':'식별 보류','excluded_passive':'비대상','needs_scope_review':'범위 확인','not_searched':'미조사'}
KIND={'source_code':'공개 소스 코드','driver':'드라이버','sdk_api':'SDK/API','programming_manual':'프로그래밍/통신 명령서','control_software':'제어 소프트웨어','product_interface':'제품/통신 사양','support_portal':'지원/발견 페이지'}
MATCH={'exact_model':'정확 모델','documented_family':'명시된 제품군','family_candidate':'제품군 후보','generic_vendor':'제조사 범용','discovery_only':'발견/참고'}
RANK={'exact_model':0,'documented_family':1,'family_candidate':2,'generic_vendor':3,'discovery_only':4}
def clean_result(value):
 if isinstance(value,str):return re.sub(r'[\x00-\x08\x0b\x0c\x0e-\x1f]',' ',value)
 if isinstance(value,list):return [clean_result(v) for v in value]
 if isinstance(value,dict):return {k:clean_result(v) for k,v in value.items()}
 return value
def code_status(sources):
 codes=[s for s in sources if s.get('kind')=='source_code']
 if not codes:return '공개 소스 코드 미확인'
 m=min(codes,key=lambda s:RANK.get(s.get('match_level'),9)).get('match_level')
 return {'exact_model':'모델명 일치 공개 코드','documented_family':'명시된 제품군 공개 코드','family_candidate':'제품군 공개 코드 후보','generic_vendor':'범용 공개 코드','discovery_only':'공개 코드 발견 후보'}.get(m,'공개 코드 적용성 확인 필요')
def main():
 inventory=json.loads((ROOT/'inventory_items.json').read_text())
 meta=json.loads((ROOT/'inventory_snapshot.json').read_text())
 base={x['item_id']:x for x in inventory}; result={}
 for p in sorted((ROOT/'results').glob('*/results.json')):
  for r in json.loads(p.read_text()):
   r=clean_result(r)
   k=r['item_id']
   if k in result:raise ValueError('duplicate item result: '+k)
   if k not in base:raise ValueError('unknown item result: '+k)
   r['research_domain']=p.parent.name
   result[k]=r
 if set(base)!=set(result):raise ValueError('missing results: '+str(len(set(base)-set(result))))
 pending=[r['item_id'] for r in result.values() if r.get('research_status')=='not_searched']
 if pending:raise ValueError('not_searched items remain: '+str(len(pending)))
 merged=[]; all_sources=[]; all_queries=[]; unique_url=set(); risks=[]
 for i in inventory:
  r=result[i['item_id']]
  for f in ['scope','scope_reason','research_status','sources','searched_queries','search_basis','notes']:
   if f not in r:raise ValueError(f'missing {f}: '+i['item_id'])
  sources=r.get('sources',[])
  if not isinstance(sources,list):raise TypeError('sources not array')
  for s in sources:
   u=s.get('url','')
   if not u.startswith(('https://','http://')):raise ValueError('invalid source URL: '+u)
   if '...' in u or '…' in u:risks.append({'item_id':i['item_id'],'url':u,'issue':'URL 축약 표기 확인 필요'})
   unique_url.add(u)
   all_sources.append(dict(item_id=i['item_id'],manufacturer=i['manufacturer'],model=i['model'],**s))
  for q in r.get('searched_queries',[]):all_queries.append({'item_id':i['item_id'],'query':q,'search_basis':r.get('search_basis','')})
  row=dict(i,**r);row['public_code_status']=code_status(sources);row['hardware_verified']=False;row['reviewed_at_kst']='2026-10-24'
  merged.append(row)
 byscope=collections.Counter(r['scope'] for r in merged)
 row_scope=collections.Counter()
 for r in merged:row_scope[r['scope']]+=r['registration_rows']
 code_exact=sum(r['scope']=='included' and any(s.get('kind')=='source_code' and s.get('match_level')=='exact_model' for s in r['sources']) for r in merged)
 code_family=sum(r['scope']=='included' and not any(s.get('kind')=='source_code' and s.get('match_level')=='exact_model' for s in r['sources']) and any(s.get('kind')=='source_code' and s.get('match_level')=='documented_family' for s in r['sources']) for r in merged)
 kind_model=collections.Counter()
 for r in merged:
  if r['scope']!='excluded_passive':
   for k in {s.get('kind') for s in r['sources']}:kind_model[k]+=1
 summary=dict(meta,completed_item_records=len(merged),unsearched_items=0,scope_counts=dict(byscope),scope_registration_rows=dict(row_scope),research_status_counts=dict(collections.Counter(r['research_status'] for r in merged)),unique_source_urls=len(unique_url),model_source_associations=len(all_sources),search_query_associations=len(all_queries),exact_public_code_items=code_exact,documented_family_public_code_items=code_family,source_kind_item_counts=dict(kind_model),hardware_verified_items=0,url_format_warnings=len(risks))
 OUT.mkdir(exist_ok=True)
 (OUT/'summary.json').write_text(json.dumps(summary,ensure_ascii=False,indent=2))
 (OUT/'model_control_sources.json').write_text(json.dumps(merged,ensure_ascii=False,indent=2))
 (OUT/'url_format_warnings.json').write_text(json.dumps(risks,ensure_ascii=False,indent=2))
 fields=['item_id','manufacturer','model','names','registration_rows','scope','scope_reason','research_status','public_code_status','search_basis','notes','hardware_verified']
 with (OUT/'model_control_sources.csv').open('w',encoding='utf-8-sig',newline='') as f:
  w=csv.DictWriter(f,fieldnames=fields+['source_urls']);w.writeheader()
  for r in merged:
   d={k:r.get(k,'') for k in fields};d['names']='; '.join(r['names']);d['source_urls']='\n'.join(s['url'] for s in r['sources']);w.writerow(d)
 dbpath=OUT/'equipment_control_research.sqlite3'
 if dbpath.exists():dbpath.unlink()
 db=sqlite3.connect(dbpath)
 db.executescript("""
 create table metadata(key text primary key,value_json text not null);
 create table models(item_id text primary key,manufacturer text,model text,names_json text,registration_rows integer,scope text,research_status text,public_code_status text,scope_reason text,notes text,search_basis text,hardware_verified integer);
 create table register_mapping(equipment_id text primary key,register_no integer,item_id text references models(item_id));
 create table sources(source_id integer primary key,url text unique not null);
 create table model_sources(item_id text references models(item_id),source_id integer references sources(source_id),kind text,authority text,match_level text,evidence text,access text,license text);
 create table searches(item_id text references models(item_id),query text,search_basis text);
 """)
 db.execute('insert into metadata values(?,?)',('summary',json.dumps(summary,ensure_ascii=False)))
 for r in merged:
  db.execute('insert into models values(?,?,?,?,?,?,?,?,?,?,?,?)',(r['item_id'],r['manufacturer'],r['model'],json.dumps(r['names'],ensure_ascii=False),r['registration_rows'],r['scope'],r['research_status'],r['public_code_status'],r['scope_reason'],r['notes'],r['search_basis'],0))
  for e,n in zip(r['equipment_ids'],r['register_numbers']):db.execute('insert into register_mapping values(?,?,?)',(e,n,r['item_id']))
 for s in all_sources:
  db.execute('insert or ignore into sources(url) values(?)',(s['url'],));sid=db.execute('select source_id from sources where url=?',(s['url'],)).fetchone()[0]
  db.execute('insert into model_sources values(?,?,?,?,?,?,?,?)',(s['item_id'],sid,s.get('kind',''),s.get('authority',''),s.get('match_level',''),s.get('evidence',''),s.get('access',''),s.get('license','미확인')))
 db.executemany('insert into searches values(?,?,?)',[(q['item_id'],q['query'],q['search_basis']) for q in all_queries])
 db.commit();assert db.execute('pragma integrity_check').fetchone()[0]=='ok';assert db.execute('select count(*) from register_mapping').fetchone()[0]==6470;db.close()
 wb=Workbook();ws=wb.active;ws.title='읽어주세요'
 intro=[
 ['전체 장비 제어자료 조사','2026-10-24 Asia/Seoul'],
 ['원본 등록행',meta['inventory_rows']],['제조사+모델 레코드',meta['identified_models']],['식별정보 미완성 원본행',meta['unidentified_rows']],['감사한 전체 항목',len(merged)],['제어·계측 대상',byscope['included']],['제조사·모델 식별 보류',byscope['identity_unresolved']],['제어대상 여부 확인',byscope['needs_scope_review']],['비대상 부품·치구',byscope['excluded_passive']],['정확 모델 공개 코드 발견',code_exact],['명시 제품군 공개 코드 발견(위 수와 중복 제외)',code_family],['고유 출처 URL',len(unique_url)],['실장비 검증', '미실행'],
 ['읽는 순서','제어대상 → 출처_링크 → 확인필요. 전체대조에서 원본 등록번호 추적 가능.'],
  ['모델명 일치 자료','모델을 명시한 소스/문서를 찾았다는 뜻. 제품/통신 사양만 있는 경우도 포함. 전체 기능·설치·실기 호환을 보증하지 않음.'],
  ['공개 코드 집계','제어·계측 대상으로 분류된 항목만 집계. 제조사 불일치 등 식별 보류 항목의 후보 코드는 제외. 제조사 원문 표기가 다르면 같은 제품도 별도 항목일 수 있음.'],
 ['공개 코드','MIT/Apache와 제품 한정 샘플 라이선스를 구분. 무료 실행파일·API는 공개 소스와 별개.'],
 ['검색 방식','개별 모델 검색 또는 해당 모델을 명시한 제품군 지원표/실제 공개 코드 파일을 대조. 검색 쿼리 수는 모델 수와 다름.'],
 ['공개자료 미확보','이번 검색으로 찾지 못한 것. 자료가 없거나 원격제어가 불가능하다는 확정 아님.'],
 ['원문 확인','search_index_only는 검색 색인 확인, content_read/download_read는 열람. blocked/redirected는 제한 또는 경로 변경.'],
 ['물리 설비','정적 차폐실·단순 지지대 자체와 제어기·마스트·모터를 분리. 원본 항목은 삭제하지 않음.'],
 ['자료 취급','원본 DB 변경 없음. 공개 코드 설치·실행 및 장비 접속/동작 명령 없음.'],
 ['데이터 파일','같은 폴더 JSON/CSV/SQLite 제공. 원시 검색 로그와 재현 스크립트는 한 단계 위 research 폴더에 보존.'],
 ['원본 DB SHA256',meta['database_sha256']]
 ]
 for row in intro:ws.append(row)
 headers=['항목ID','제조사(원문)','모델(원문)','대장 장비명','등록행 수','범위','조사 결과','공개 코드 상태','대표 공개 코드 링크','대표 제어자료 링크','제외·보류 근거','추가 확인','실장비 검증']
 def pick(ss,kinds):
  subset=sorted([s for s in ss if s.get('kind') in kinds],key=lambda s:RANK.get(s.get('match_level'),9))
  return subset[0]['url'] if subset else ''
 def mr(r):
  return [r['item_id'],r['manufacturer'],r['model'],'; '.join(r['names']),r['registration_rows'],SCOPE.get(r['scope'],r['scope']),STATUS.get(r['research_status'],r['research_status']),r['public_code_status'],pick(r['sources'],{'source_code'}),pick(r['sources'],{'driver','sdk_api','programming_manual','control_software','product_interface'}),r['scope_reason'],r['notes'],'미실행']
 ordered=sorted(merged,key=lambda r:(r['manufacturer'].casefold(),r['model'].casefold(),r['item_id']))
 for title,subset in [('제어대상',[r for r in ordered if r['scope']=='included']),('확인필요',[r for r in ordered if r['scope'] in ('identity_unresolved','needs_scope_review')]),('제외항목',[r for r in ordered if r['scope']=='excluded_passive'])]:
  w=wb.create_sheet(title);w.append(headers)
  for r in subset:w.append(mr(r))
 w=wb.create_sheet('출처_링크');w.append(['항목ID','제조사','모델','종류','공식·커뮤니티 구분','모델 대응 수준','링크','근거','열람 수준','라이선스'])
 for s in all_sources:w.append([s['item_id'],s['manufacturer'],s['model'],KIND.get(s.get('kind'),s.get('kind')),s.get('authority'),MATCH.get(s.get('match_level'),s.get('match_level')),s['url'],s.get('evidence'),s.get('access'),s.get('license')])
 w=wb.create_sheet('검색기록');w.append(['항목ID','실제 실행 검색어','검색/파일 검토 근거'])
 for q in all_queries:w.append([q['item_id'],q['query'],q['search_basis']])
 w=wb.create_sheet('전체대조');w.append(['대장 등록번호','장비ID','항목ID','제조사','모델','대장 장비명','최종 범위','조사 결과'])
 rows=[]
 for r in merged:
  for e,n in zip(r['equipment_ids'],r['register_numbers']):rows.append([n,e,r['item_id'],r['manufacturer'],r['model'],'; '.join(r['names']),SCOPE.get(r['scope'],r['scope']),STATUS.get(r['research_status'],r['research_status'])])
 for row in sorted(rows,key=lambda a:a[0]):w.append(row)
 w=wb.create_sheet('분류복원');w.append(headers)
 for r in ordered:
  if r['initial_scope']=='excluded_passive' and r['scope']!='excluded_passive':w.append(mr(r))
 for idx,sheet in enumerate(wb.worksheets):
  sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
  for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='174B67');cell.alignment=Alignment(wrap_text=True,vertical='top')
  sheet.row_dimensions[1].height=30
  for col in sheet.columns:
   letter=col[0].column_letter;sheet.column_dimensions[letter].width=min(52,max(16,len(str(col[0].value or ''))*2+4))
  for row in sheet.iter_rows(min_row=2):
   for cell in row:
    if isinstance(cell.value,str):
     # Treat all source/register text as data rather than spreadsheet formulas.
     if cell.value.startswith(('=','+','-','@')):cell.data_type='s'
     if cell.value.startswith(('https://','http://')):cell.hyperlink=cell.value;cell.font=Font(color='0563C1',underline='single')
     if len(cell.value)>32767:cell.value=cell.value[:32700]+' [전체 내용은 JSON 참조]'
    cell.alignment=Alignment(vertical='top',wrap_text=True)
  if sheet.title=='읽어주세요':sheet.column_dimensions['A'].width=48;sheet.column_dimensions['B'].width=105
 wb.save(OUT/'전체장비_제어자료_조사.xlsx')
 check=load_workbook(OUT/'전체장비_제어자료_조사.xlsx',read_only=True)
 assert check['전체대조'].max_row-1==6470
 assert sum(check[n].max_row-1 for n in ['제어대상','확인필요','제외항목'])==len(merged)
 check.close()
 (OUT/'validation.json').write_text(json.dumps({'all_inventory_rows_accounted':True,'all_item_ids_accounted_once':True,'unsearched_items':0,'sqlite_integrity':'ok','excel_rows_validated':True,'hardware_tested':False},ensure_ascii=False,indent=2))
 print(json.dumps(summary,ensure_ascii=False,indent=2))
if __name__=='__main__':main()

